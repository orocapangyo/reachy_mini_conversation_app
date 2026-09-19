import docx


def dump_doc(path):
    print("=== DUMP FOR:", path, "===")
    doc = docx.Document(path)
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"P{p_idx}: {p.text.strip()}")
    for i, t in enumerate(doc.tables):
        print(f"--- Table {i} ({len(t.rows)} rows, {len(t.columns)} cols) ---")
        for r_idx, row in enumerate(t.rows):
            cells_text = [f"c{c_idx}: {c.text.strip()}" for c_idx, c in enumerate(row.cells)]
            print(f"  Row {r_idx}: " + " | ".join(cells_text))


if __name__ == "__main__":
    dump_doc(r"docs/oss_report/(해당시 제출) 출품작 중복수혜 여부 확인서_접수번호(팀명).docx")
    print("\n" + "=" * 60 + "\n")
    dump_doc(r"docs/oss_report/2026 오픈소스 개발자대회 결과보고서_접수번호(팀명).docx")
